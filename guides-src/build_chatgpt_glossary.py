"""Build assets/guides/Blach_ChatGPT_AI_Glossary.docx from the Claude glossary.

The Claude glossary is the formatting template: its title, subtitle, rule,
intro, section-heading, term and note paragraphs are cloned and refilled, so
both glossaries share one look. Feature names are current as of October 2026.

Run:  python guides-src/build_chatgpt_glossary.py
"""
import copy
import datetime
import pathlib

import docx

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "assets/guides/Blach_Claude_AI_Glossary.docx"
OUT = ROOT / "assets/guides/Blach_ChatGPT_AI_Glossary.docx"

TITLE = "ChatGPT AI Glossary"
SUBTITLE = "A shared vocabulary for the Blach × ChatGPT rollout"
INTRO = ("This is a quick desk reference for the terms most likely to come up as you use ChatGPT at Blach. "
         "Definitions are plain-language by design — no technical background needed to follow along. "
         "ChatGPT changes quickly; feature names here are current as of October 2026.")

# (heading, optional italic note, [(term, definition), ...])
SECTIONS = [
    ("AI & ChatGPT Foundations", None, [
        ("OpenAI", "The company that builds ChatGPT. Blach provides ChatGPT through a company-managed OpenAI workspace."),
        ("ChatGPT", "One of the two AI assistants Blach provides, alongside Claude. It reads, writes, analyzes, and — in Work mode — carries out multi-step tasks on your behalf."),
        ("Model", "The engine behind each answer. Blach’s ChatGPT models come in three tiers: Sol (most capable), Terra (balanced, for everyday work) and Luna (fastest). You pick one from the model menu."),
        ("Reasoning effort", "How long the model thinks before it answers, from Instant up to Extra High. More effort gives more careful answers but uses more of your limit."),
        ("Large Language Model (LLM)", "The kind of AI ChatGPT is: a system trained on vast amounts of text that predicts and generates language. It’s the engine behind everything ChatGPT does."),
        ("Generative AI", "AI that creates new content — drafts, summaries, images — rather than only retrieving existing answers."),
        ("Prompt", "The message or instructions you give ChatGPT. Clearer, more specific prompts get better results."),
        ("Token", "The small chunk of text (roughly three-quarters of a word; 100 tokens is about 75 words) that ChatGPT reads and writes in. Usage and limits are measured in tokens."),
        ("Context window", "How much information ChatGPT can hold in view at once — the current conversation, attached files, and instructions. It’s large, but finite."),
        ("Knowledge cutoff", "The date after which a model wasn’t trained on new information. For anything more recent, ChatGPT searches the web."),
        ("Hallucination", "When ChatGPT states something that sounds confident but is wrong. This is why a person should review important outputs before they go out."),
    ]),
    ("How We Access ChatGPT", None, [
        ("Workspace", "Blach’s company ChatGPT account. Click the Blach logo at the bottom left to reach Settings. Administrators decide which models, features and apps are turned on."),
        ("Chat", "The familiar conversational interface (chatgpt.com, the desktop app, and the mobile app). You type, ChatGPT responds — best for questions, drafting, and back-and-forth."),
        ("Desktop app", "ChatGPT installed on your computer, one click from your taskbar. It works alongside the files and apps on your machine — best for longer, hands-on work."),
        ("Work", "ChatGPT’s agent mode for longer jobs. It plans and carries out multi-step work — research, analysis, spreadsheets, presentations, reports — using your connected apps, and checks with you before important actions. It replaced the earlier “agent mode.”"),
        ("Deep research", "A multi-step research job that reads many sources and returns a cited report. Start it from the + menu; you can adjust the research plan before it runs."),
        ("Project", "A shared folder that keeps related chats, files, and instructions together so every chat inside it starts with the right context."),
        ("Space / Page", "Space is your home for pages and files. A page is an editable, shareable document — a brief, meeting notes, a team update — that teammates with access can edit live. A space groups related pages for a topic or team."),
        ("Site", "A web page or lightweight app — a tracker, dashboard, or portal — that ChatGPT builds and hosts at a link. New sites are private until you share them."),
        ("Writing block", "An editable draft that appears right in the chat. It replaced Canvas, which you may still see in older tutorials."),
        ("ChatGPT for Excel, PowerPoint & Word", "A ChatGPT sidebar that works directly inside Office files, where enabled for our workspace."),
    ]),
    ("Customizing ChatGPT for Blach", None, [
        ("Agent / Agentic", "Describes ChatGPT when it plans and carries out multi-step tasks on its own, rather than answering one prompt at a time. Work and Codex are agentic."),
        ("Automation / Workflow", "A repeatable process handed to ChatGPT so it runs with less manual effort — for example, turning meeting notes into a summary and an action list."),
        ("Skill", "A saved set of instructions for one repeatable job that ChatGPT applies when the task fits.  e.g., our Blach branding skill applies Blach fonts, colors, and logos; Blach voice writes in our tone."),
        ("App (Connector)", "A secure link between ChatGPT and an outside system (such as SharePoint or Outlook) that lets ChatGPT look things up and work within it. OpenAI now calls connectors “apps.”"),
        ("Plugin", "A package that bundles skills and apps for a workflow, installed from the plugin directory."),
        ("Model Context Protocol (MCP)", "The open standard apps are built on. Blach’s custom connectors use it, which is why the same tools can connect to both Claude and ChatGPT."),
        ("Custom GPT", "A customized ChatGPT with its own instructions and files. OpenAI is retiring custom GPTs on December 11, 2026 — use Projects and Skills instead."),
        ("Custom instructions", "Standing guidance you write about yourself and how you want answers, under Settings → Personalization.  e.g., “I’m a project engineer at a commercial GC. Use CSI divisions. Keep answers short.”"),
        ("Memory", "ChatGPT’s ability to carry relevant details forward across chats, so you don’t have to re-explain each time. You can view, edit, or turn it off in Settings."),
        ("Scheduled task", "A job ChatGPT runs later or on repeat — such as a weekly summary — managed from the Scheduled page in the sidebar."),
        ("Prompt engineering", "The practice of writing clear, well-structured prompts and instructions to get reliable, high-quality results."),
        ("Human-in-the-loop", "Keeping a person in charge of reviewing and approving ChatGPT’s work at key steps — our default for anything client-facing."),
    ]),
    ("Trust, Security & Confidentiality", None, [
        ("Confidentiality / Data handling", "Our rule that client names, project details, pricing, drawings, and personnel information stay protected and inside Blach. Company information stays inside Claude and ChatGPT — follow the AI Policy."),
        ("Usage limits", "Each person has their own allowance. If you hit a limit, wait until it resets — other models may still work. With credits turned off, hitting a limit won’t trigger extra charges."),
        ("Temporary chat", "A chat that stays out of your history and doesn’t use or update memory. OpenAI may still keep it for up to 30 days for safety."),
        ("Prompt injection", "A security risk where hidden instructions buried in a document, email, or web page try to trick ChatGPT into doing something it shouldn’t — a reason to be careful about what untrusted content an agent is allowed to read."),
        ("Lethal trifecta", "Shorthand for the risky combination that makes prompt injection dangerous: access to private data, plus exposure to untrusted content, plus the ability to send information out. Keeping these three from lining up is a core safety principle."),
        ("Admin controls", "Workspace settings that let our administrators manage access, turn features and apps on or off, and set guardrails for how ChatGPT is used."),
    ]),
    ("Blach Systems You’ll Hear Mentioned",
     "Blach is bringing the same connectors Claude uses over to ChatGPT. Availability changes as apps are approved, so check with the AI team before relying on one.", [
        ("Microsoft 365", "Outlook, Teams, SharePoint, and the Office apps we use every day."),
        ("Egnyte", "Our cloud file storage and digital asset management (DAM) system, where project and marketing files live."),
        ("CMiC", "Our construction enterprise resource planning (ERP) and accounting system."),
        ("Autodesk Build", "Autodesk’s construction platform for RFIs, submittals, issues, sheets, and photos."),
        ("Primavera P6", "The scheduling software used to build and track project schedules."),
        ("Clearstory", "Field-to-office platform for tracking, documenting, and pricing change events in real time."),
        ("Bridgit Bench", "Workforce and resource planning software."),
        ("HammerTech", "Our jobsite safety platform for orientations, inspections, and incident reporting."),
    ]),
]


def set_runs(p_el, texts):
    """Write texts into the paragraph's runs in order (run formatting is kept)."""
    runs = p_el.findall(docx.oxml.ns.qn("w:r"))
    assert len(runs) >= len(texts), (len(runs), texts)
    for r, t in zip(runs, texts):
        r.find(docx.oxml.ns.qn("w:t")).text = t
    for r in runs[len(texts):]:
        p_el.remove(r)
    return p_el


def main():
    d = docx.Document(str(SRC))
    paras = d.paragraphs
    assert paras[0].text == "Claude AI Glossary", "template changed — re-check paragraph indices"
    heading_t = copy.deepcopy(paras[4]._p)
    term_t = copy.deepcopy(paras[5]._p)
    note_t = copy.deepcopy(paras[41]._p)

    set_runs(paras[0]._p, [TITLE])
    set_runs(paras[1]._p, [SUBTITLE])
    set_runs(paras[3]._p, [INTRO])

    body = paras[0]._p.getparent()
    for p in paras[4:]:
        body.remove(p._p)
    sect = body.find(docx.oxml.ns.qn("w:sectPr"))

    def add(el):
        sect.addprevious(el)

    for i, (heading, note, terms) in enumerate(SECTIONS, 1):
        add(set_runs(copy.deepcopy(heading_t), ["%02d  " % i, heading]))
        if note:
            add(set_runs(copy.deepcopy(note_t), [note]))
        for term, definition in terms:
            add(set_runs(copy.deepcopy(term_t), [term + "  ", "— " + definition]))

    for para in d.sections[0].footer.paragraphs:
        for r in para.runs:
            r.text = r.text.replace("Claude AI Glossary", TITLE)

    cp = d.core_properties
    cp.title = TITLE
    cp.description = "Shared vocabulary for the Blach x ChatGPT rollout"
    cp.created = cp.modified = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
    d.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
